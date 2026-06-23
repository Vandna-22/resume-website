import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def create_cv(output_path='Vandna_Gupta_CV_Generated.docx'):
    doc = docx.Document()

    # Set page margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    DARK_BLUE = RGBColor(31, 78, 121)
    BLACK = RGBColor(0, 0, 0)
    DARK_GRAY = RGBColor(64, 64, 64)

    # Header - Centered Name
    p_name = doc.add_paragraph()
    p_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_name = p_name.add_run('Vandna Gupta')
    r_name.bold = True
    r_name.font.size = Pt(28)
    r_name.font.name = 'Times New Roman'
    p_name.paragraph_format.space_after = Pt(0)

    # Header - Centered Location
    p_loc = doc.add_paragraph()
    p_loc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_loc = p_loc.add_run('Khurja, Uttar Pradesh, India | Website: vandna.manupal.dev')
    r_loc.font.size = Pt(11)
    r_loc.font.name = 'Arial'
    r_loc.font.color.rgb = DARK_GRAY
    p_loc.paragraph_format.space_after = Pt(4)

    # Header - Contact Info
    p_contact = doc.add_paragraph()
    p_contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_contact.paragraph_format.space_after = Pt(16)
    
    def add_contact_part(label, value):
        r1 = p_contact.add_run(label)
        r1.bold = True
        r1.font.name = 'Arial'
        r1.font.size = Pt(10)
        
        r2 = p_contact.add_run(value)
        r2.font.name = 'Arial'
        r2.font.size = Pt(10)
        r2.font.color.rgb = DARK_BLUE
        
    add_contact_part('Email: ', 'vandnaa2230@gmail.com     ')
    add_contact_part('LinkedIn: ', 'linkedin.com/in/vandna-gupta-6a579a351     ')
    add_contact_part('GitHub: ', 'github.com/Vandna-22')

    # Helper for Headings
    def add_heading(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(6)
        r = p.add_run(text)
        r.bold = True
        r.font.size = Pt(16)
        r.font.name = 'Times New Roman'
        r.font.color.rgb = DARK_BLUE
        
        # Add bottom border
        pBdr = OxmlElement('w:pBdr')
        bottom = OxmlElement('w:bottom')
        bottom.set(qn('w:val'), 'single')
        bottom.set(qn('w:sz'), '6')
        bottom.set(qn('w:space'), '1')
        bottom.set(qn('w:color'), '1F4E79')
        pBdr.append(bottom)
        p._p.get_or_add_pPr().append(pBdr)
        
        return p

    def add_normal_text(text, bold=False, italic=False, color=BLACK):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(text)
        r.font.name = 'Arial'
        r.font.size = Pt(10)
        r.bold = bold
        r.italic = italic
        r.font.color.rgb = color
        return p

    def add_bullet(text):
        p = doc.add_paragraph(text, style='List Bullet')
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.left_indent = Inches(0.25)
        for r in p.runs:
            r.font.name = 'Arial'
            r.font.size = Pt(10)
        return p

    # --- Professional Profile ---
    add_heading('Professional Profile')
    p = doc.add_paragraph()
    r = p.add_run('M.Sc. Bioinformatics graduate (Chaudhary Charan Singh University, Completed with 86.25%) with an active research internship at ICMR-National Institute of Cancer Prevention & Research (NICPR), Noida. Specializes in end-to-end metatranscriptomics pipeline development, functional biomarker profiling using the KEGG database, NGS-based microbiome analysis, and multi-omics data interpretation in oral cancer research. Distinguished by advanced proficiency in frontier AI research tools — Claude AI (Anthropic), NeodiSC, AntiGravity AI, OpenCode AI, and Gemini CLI — enabling accelerated hypothesis generation, literature synthesis, and bioinformatics workflow design at a level uncommon among early-career researchers. Also experienced in molecular docking and protein-ligand simulation using AutoDock and AutoDock Vina. Certified in Advanced Microsoft Excel.')
    r.font.name = 'Arial'
    r.font.size = Pt(10)
    p.paragraph_format.space_after = Pt(10)

    # --- Key Achievements ---
    add_heading('Key Achievements')
    add_bullet('University Gold Medalist in M.Sc. Bioinformatics (86.25%), Chaudhary Charan Singh University')
    add_bullet('Co-authored a systematic review and meta-analysis on oral microbiome biomarkers in OSCC — manuscript submitted to Elsevier')
    add_bullet('Built a complete end-to-end metatranscriptomics pipeline (SRA Toolkit → FastQC → Trimmomatic → Bowtie2 → Kraken2/Bracken → STAR/HISAT2 → DESeq2/edgeR → KEGG) from scratch during internship')
    add_bullet('One of very few early-career bioinformaticians integrating frontier AI tools (Claude AI, Gemini CLI, OpenCode) into active cancer genomics research workflows')
    add_bullet('Developed an automated Python pipeline for extraction of BioSample IDs from research literature PDFs to streamline data collection for meta-analysis studies')
    add_bullet('Certified in Advanced Microsoft Excel (2025) for data analysis and reporting')

    # --- Education ---
    add_heading('Education')
    
    table_edu = doc.add_table(rows=2, cols=3)
    table_edu.autofit = False
    table_edu.columns[0].width = Inches(2.0)
    table_edu.columns[1].width = Inches(3.0)
    table_edu.columns[2].width = Inches(1.5)

    def fill_row(row, col0, col1, col2, bold=True):
        row.cells[0].text = col0
        row.cells[1].text = col1
        row.cells[2].text = col2
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(2)
                for r in p.runs:
                    r.font.name = 'Arial'
                    r.font.size = Pt(10)
                    r.bold = bold

    fill_row(table_edu.rows[0], 'Completed (86.25%)', 'M.Sc. Bioinformatics (Gold Medalist)', 'CCSU, Meerut')
    
    p_course = doc.add_paragraph()
    r_c1 = p_course.add_run('Relevant Coursework: ')
    r_c1.bold = True
    r_c1.font.name = 'Arial'
    r_c1.font.size = Pt(10)
    r_c2 = p_course.add_run('Genomics, Proteomics, Bioinformatics Algorithms, Molecular Biology, Computational Biology, Biostatistics')
    r_c2.font.name = 'Arial'
    r_c2.font.size = Pt(10)
    p_course.paragraph_format.space_after = Pt(10)

    fill_row(table_edu.rows[1], 'Completed', 'B.Sc. Biotechnology', 'CCSU, Meerut')
    table_edu.rows[1].cells[0].paragraphs[0].runs[0].bold = True

    # --- Research & Internship Experience ---
    add_heading('Research & Internship Experience')
    
    table_exp1 = doc.add_table(rows=1, cols=3)
    table_exp1.autofit = False
    table_exp1.columns[0].width = Inches(2.0)
    table_exp1.columns[1].width = Inches(3.0)
    table_exp1.columns[2].width = Inches(1.5)
    fill_row(table_exp1.rows[0], '2025 – Present', 'Bioinformatics Research Trainee', 'ICMR-NICPR, Noida')

    add_normal_text('Supervisor: Dr. Pramod Kumar (Scientist-D), Division of Molecular Biology', italic=True, color=RGBColor(102, 102, 102))
    add_normal_text('Main responsibilities and contributions:', bold=True)
    
    bullets_exp1 = [
        'Designed and executed a complete metatranscriptomics pipeline for microbial community profiling from NCBI BioProject sequencing data, covering SRA retrieval, QC, host removal, taxonomic classification, RNA-seq alignment, differential expression, and KEGG functional interpretation',
        'Conducted functional profiling of cancer-associated biomarkers using the KEGG database to map metabolic and signaling pathways implicated in OSCC progression',
        'Analyzed NGS-derived microbiome datasets from oral and other cancer cohorts using Linux-based bioinformatics workflows in R and Python',
        'Applied statistical methods including Mann-Whitney and Kruskal-Wallis tests for differential abundance testing and multi-group comparisons',
        'Leveraged Claude AI (Anthropic) for AI-assisted literature review, rapid hypothesis generation, research documentation, and manuscript preparation',
        'Utilized NeodiSC and AntiGravity AI tools for specialized bioinformatics analysis and biomarker discovery workflows',
        'Generated publication-quality data visualizations — violin plots, boxplots, heatmaps — using ggplot2 and ComplexHeatmap in R',
        'Developed Python automation tools: PDF classifier (smoking vs. non-smoking papers), BioSample ID extraction pipeline from literature PDFs, AI-powered PubMed screening app for systematic review triage'
    ]
    for b in bullets_exp1:
        add_bullet(b)

    doc.add_paragraph() # space

    table_exp2 = doc.add_table(rows=1, cols=3)
    table_exp2.autofit = False
    table_exp2.columns[0].width = Inches(2.0)
    table_exp2.columns[1].width = Inches(3.0)
    table_exp2.columns[2].width = Inches(1.5)
    fill_row(table_exp2.rows[0], '2023 (May – Aug)', 'Summer Research Trainee', 'Bebymil International Pvt. Ltd.')
    add_bullet('Carried out summer research training on "Basic Techniques of Molecular and Microbiology"')
    add_bullet('Gained practical laboratory experience in core molecular biology protocols and microbiological techniques under professional supervision')

    # --- Research Projects ---
    add_heading('Research Projects')
    proj_data = [
        ('Oral Microbiome Meta-Analysis in OSCC — Systematic Review (Manuscript Submitted)', [
            'Co-authored systematic review analyzing oral microbiome composition across global OSCC patient cohorts; identified conserved biomarker taxa and assessed confounding impact of age and geographic variation',
            'AI-accelerated synthesis of 50+ peer-reviewed studies using Claude AI and OpenCode AI'
        ]),
        ('Literature Mining Pipeline for Oral Host Microbiome', [
            'Developed an automated pipeline for targeted paper searching on meta-transcriptomics in the oral host microbiome.',
            'Expanded the literature extraction framework to systematically identify and categorize papers related to OPMD and OSCC within metagenomic studies.'
        ]),
        ('Metagenomics Pipeline — 16S rRNA Amplicon Analysis (QIIME2 + DADA2)', [
            'Processed raw 16S rRNA amplicon sequencing data through a complete QIIME2 pipeline: paired-end import → DADA2 denoising → ASV table generation → taxonomic classification using SILVA database',
            'Generated diversity analyses (alpha and beta diversity), taxonomic bar plots (phylum to species level), and differential abundance outputs to characterize oral microbial community shifts in OSCC vs. healthy controls',
            'Produced publication-quality taxonomic bar plots and PCoA plots to visualize community composition across patient cohorts'
        ]),
        ('Functional Biomarker Profiling via KEGG Database', [
            'Mapped cancer-associated microbial taxa to KEGG metabolic and signaling pathways using DADA2 ASV pipeline, 16S rRNA profiling, and QIIME2',
            'Integrated Claude AI-assisted annotation with statistical outputs for accelerated pathway-level biological interpretation'
        ]),
        ('Molecular Docking & Simulation — AutoDock / AutoDock Vina', [
            'Performed protein-ligand molecular docking using AutoDock and AutoDock Vina to evaluate binding affinities and interaction modes of candidate bioactive compounds against cancer-relevant target proteins',
            'Prepared receptor and ligand structures (PDB processing, energy minimization), defined grid boxes, and analyzed docking poses using PyMOL and Discovery Studio Visualizer',
            'Applied GROMACS and NAMD for molecular dynamics simulations to assess stability of docked complexes over simulation trajectories'
        ]),
        ('Metabolomics ML Classification Pipeline — Control vs. OSCC', [
            'Developed ML classification pipeline using XGBoost, Random Forest, SVM, and Neural Networks (scikit-learn, TensorFlow) for metabolomics data analysis'
        ]),
        ('Microbiome Statistical Analysis & Visualization', [
            'Performed differential abundance testing of microbial taxa in OSCC vs. healthy controls; reproduced diverging orange-blue heatmap of cellular biomarkers in R (ComplexHeatmap) and Python (seaborn/matplotlib)'
        ])
    ]
    for title, bullets in proj_data:
        add_normal_text(title, bold=True)
        for b in bullets:
            add_bullet(b)

    # --- Technical Skills ---
    add_heading('Technical Skills')
    skills_data = [
        ('OS & Shell', 'Linux (Ubuntu) — command line, shell scripting, pipeline execution'),
        ('Languages', 'Python, R, Bash'),
        ('Metatranscriptomics', 'SRA Toolkit, FastQC, Trimmomatic, Bowtie2, Kraken2, Bracken, STAR, HISAT2, DESeq2, edgeR'),
        ('Metagenomics', 'QIIME2, DADA2, 16S rRNA amplicon analysis, taxonomic bar plots, alpha/beta diversity, SILVA database'),
        ('Molecular Docking', 'AutoDock, AutoDock Vina, PyMOL, Discovery Studio Visualizer — protein-ligand docking & binding affinity analysis'),
        ('MD Simulations', 'GROMACS, NAMD — molecular dynamics simulation, trajectory analysis, complex stability assessment'),
        ('Machine Learning', 'XGBoost, Gradient Boosting, Random Forest, SVM, Neural Networks, Scikit-learn, TensorFlow'),
        ('AI Research Tools', 'Claude AI (Anthropic), NeodiSC, AntiGravity AI, OpenCode AI, Gemini CLI'),
        ('Databases', 'KEGG, NCBI, BioProject, ClinVar, SILVA, PDB, genomic databases'),
        ('Visualisation', 'ggplot2, ComplexHeatmap, seaborn, matplotlib, QIIME2 plots, PyMOL'),
        ('Data & Reporting', 'Advanced Microsoft Excel (Certified), Jupyter Notebook, RStudio'),
        ('Statistics', 'Mann-Whitney, Kruskal-Wallis, differential expression analysis, meta-analysis')
    ]
    table_skills = doc.add_table(rows=len(skills_data), cols=2)
    table_skills.autofit = False
    table_skills.columns[0].width = Inches(2.0)
    table_skills.columns[1].width = Inches(4.5)
    
    for i, (k, v) in enumerate(skills_data):
        c1 = table_skills.cell(i, 0)
        c2 = table_skills.cell(i, 1)
        c1.text = k
        c2.text = v
        
        for p in c1.paragraphs:
            p.paragraph_format.space_after = Pt(4)
            for r in p.runs:
                r.bold = True
                r.font.name = 'Arial'
                r.font.size = Pt(10)
                
        for p in c2.paragraphs:
            p.paragraph_format.space_after = Pt(4)
            for r in p.runs:
                r.font.name = 'Arial'
                r.font.size = Pt(10)

    # --- Certifications & Training ---
    add_heading('Certifications & Training')
    table_certs = doc.add_table(rows=3, cols=3)
    table_certs.autofit = False
    table_certs.columns[0].width = Inches(2.0)
    table_certs.columns[1].width = Inches(3.0)
    table_certs.columns[2].width = Inches(1.5)

    fill_row(table_certs.rows[0], '2025', 'Advanced Microsoft Excel', 'E-Certificate')
    fill_row(table_certs.rows[1], '2025–Present', 'Bioinformatics Research Internship', 'ICMR-NICPR, Noida')
    fill_row(table_certs.rows[2], '2023', 'Basic Techniques of Molecular and Microbiology', 'Bebymil International Pvt. Ltd.')

    doc.save(output_path)
    print(f"CV successfully generated at {output_path}")

if __name__ == '__main__':
    create_cv()
