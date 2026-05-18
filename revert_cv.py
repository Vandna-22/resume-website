import docx

doc = docx.Document('/home/vandna/Resume/Vandna_Gupta_CV_Classic(1).docx')

# Revert Paragraph 6
doc.paragraphs[6].text = "M.Sc. Bioinformatics candidate (Chaudhary Charan Singh University, 2024–2026) with an active research internship at ICMR-National Institute of Cancer Prevention & Research (NICPR), Noida. Specializes in end-to-end metatranscriptomics pipeline development, functional biomarker profiling using the KEGG database, NGS-based microbiome analysis, and multi-omics data interpretation in oral cancer research. Distinguished by advanced proficiency in frontier AI research tools — Claude AI (Anthropic), NeodiSC, AntiGravity AI, OpenCode AI, and Gemini CLI — enabling accelerated hypothesis generation, literature synthesis, and bioinformatics workflow design at a level uncommon among early-career researchers. Also experienced in molecular docking and protein-ligand simulation using AutoDock and AutoDock Vina. Certified in Advanced Microsoft Excel."

# Revert Table
table = doc.tables[0]
for row in table.rows:
    if "Databases & Analytics" in row.cells[0].text:
        row.cells[0].text = "Databases"
        row.cells[1].text = "KEGG, NCBI, BioProject, ClinVar, SILVA, PDB, genomic databases"

# Remove the Proteogenomics row. We know it was added at the very end.
last_row = table.rows[-1]
if "Proteogenomics" in last_row.cells[0].text:
    # docx python doesn't easily delete table rows, but we can access the underlying XML
    tbl = table._tbl
    tr = last_row._tr
    tbl.remove(tr)

doc.save('/home/vandna/Resume/Vandna_Gupta_CV_Classic(1).docx')
print("CV reverted successfully")
