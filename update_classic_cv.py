import docx

doc = docx.Document('/home/vandna/Resume/Vandna_Gupta_CV_Classic(1).docx')

# Update Professional Summary
doc.paragraphs[6].text = "M.Sc. Bioinformatics graduate (Chaudhary Charan Singh University, Completed with 86.25%) with active research internship at the National Institute of Cancer Prevention and Research (ICMR-NICPR), Noida. Specializes in designing and executing reproducible analysis pipelines for diverse genomic data types (WGS, WEX, GWAS) and managing genetic data workflows in internal and trusted research environments (TREs).\n\nDistinguished by advanced proficiency in designing agentic AI workflows for automated genetic evidence generation and literature synthesis. Skilled in integrating reference databases (OpenTargets, ClinVar, GWAS Catalog) with bioinformatics pipelines (Nextflow, Snakemake) and leveraging containerized environments (Docker) to accelerate computational biology research."

# Update the table
table = doc.tables[0]
for row in table.rows:
    if "Databases" in row.cells[0].text:
        row.cells[0].text = "Databases & Analytics"
        row.cells[1].text = "OpenTargets, GWAS Catalog, ClinVar, dbSNP, OMIM, ChEMBL, DrugBank, KEGG, NCBI BioProject"

# Add proteogenomics
new_row = table.add_row()
new_row.cells[0].text = "Proteogenomics"
new_row.cells[1].text = "Proteogenomic pipeline"

doc.save('/home/vandna/Resume/Vandna_Gupta_CV_Classic(1).docx')
print("CV updated successfully")
